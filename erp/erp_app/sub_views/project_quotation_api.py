import logging
from decimal import Decimal
from io import BytesIO

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError
from django.http import HttpResponse, JsonResponse
from django.views.decorators.csrf import csrf_protect
from rest_framework import status
from rest_framework.decorators import api_view
from rest_framework.exceptions import ValidationError as DRFValidationError
from rest_framework.response import Response

from ..project_quotation_items_serializer import ProjectQuotationItemSerializer
from ..project_quotation_summary_serializer import ProjectQuotationSummarySerializer
from ..sub_models.CostType_mod import CostTypeInfo
from ..sub_models.lab_furniture_item import LabFurnitureItem
from ..sub_models.project_quotation_items_mod import ProjectQuotationItemInfo, validate_project_quotation_hierarchy
from ..sub_models.project_quotation_summary_mod import ProjectQuotationSummaryInfo
from ..sub_models.room_data_mod import RoomDataInfo
from ..utils import normalize_text
from .project_quotation_view import ProjectQuotationItemView, ProjectQuotationSummaryView


BOM_HIERARCHY_ERROR = "Invalid BOM hierarchy: children must be linked to immediate parent level."
logger = logging.getLogger(__name__)


def _ensure_authenticated(request):
    if not request.user.is_authenticated:
        return JsonResponse({"status": "error", "message": "Authentication required."}, status=401)
    return None


def _serializer_error_response(serializer):
    errors = serializer.errors
    message = "Validation failed."
    response_status = "error"
    non_field_errors = errors.get("non_field_errors") if isinstance(errors, dict) else None
    if non_field_errors:
        message = str(non_field_errors[0])
    else:
        for value in errors.values() if isinstance(errors, dict) else []:
            if isinstance(value, (list, tuple)) and value:
                message = str(value[0])
                break
            if isinstance(value, dict):
                nested_values = list(value.values())
                if nested_values:
                    nested = nested_values[0]
                    if isinstance(nested, (list, tuple)) and nested:
                        message = str(nested[0])
                        break
    lowered = message.lower()
    if "requested qty cannot be greater than purchase qty" in lowered:
        response_status = "warning"
    response_payload = {"status": response_status, "message": message, "errors": errors}
    if response_status == "warning":
        response_payload["requires_confirmation"] = True
    return Response(response_payload, status=status.HTTP_400_BAD_REQUEST)


def _as_bool(value):
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "y", "on"}


def _item_serializer_context(request):
    return {
        "request": request,
        "allow_requested_qty_override": _as_bool(
            request.data.get("confirm_requested_qty_override")
        )
        if hasattr(request, "data")
        else False,
    }


def _to_decimal_or_none(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).strip())
    except Exception:
        return None


def _as_excel_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _resolve_import_column_map(first_row_values):
    headers = [normalize_text(value).lower().replace("_", " ") for value in first_row_values]
    header_map = {header: idx for idx, header in enumerate(headers) if header}

    if "item code" in header_map:
        return {
            "has_header": True,
            "item_category": header_map.get("item category"),
            "item_name": header_map.get("item name"),
            "item_code": header_map.get("item code"),
            "requested_qty": header_map.get("requested qty"),
            "room_name": header_map.get("room name"),
            "actual_cost": header_map.get("actual cost"),
        }

    # Fallback to the on-screen table order:
    # Room Name, Item Category, Item Name, Item Code, Item Type, Purchase Qty, Requested Qty, ...
    return {
        "has_header": False,
        "item_category": 1,
        "item_name": 2,
        "item_code": 3,
        "requested_qty": 6,
        "room_name": 0,
        "actual_cost": None,
    }


def _quotation_items_import_template_bytes():
    try:
        import openpyxl
    except ImportError:
        return None

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Quotation Items Import"
    worksheet.append(["Room Name", "Item Category", "Item Name", "Item Code", "Requested Qty"])
    worksheet.append(["Laboratory", "Category A", "Sample Item", "ITEM-001", "1"])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _safe_serializer_save(serializer, fallback_message):
    try:
        return serializer.save(), None
    except (DRFValidationError, DjangoValidationError) as exc:
        if hasattr(exc, "detail"):
            errors = exc.detail
        elif hasattr(exc, "message_dict"):
            errors = exc.message_dict
        else:
            errors = {"non_field_errors": exc.messages if hasattr(exc, "messages") else [str(exc)]}

        message = "Validation failed."
        if isinstance(errors, dict):
            for value in errors.values():
                if isinstance(value, (list, tuple)) and value:
                    message = str(value[0])
                    break
                if isinstance(value, dict):
                    nested_values = list(value.values())
                    if nested_values and isinstance(nested_values[0], (list, tuple)) and nested_values[0]:
                        message = str(nested_values[0][0])
                        break
        lowered = message.lower()
        response_status = "warning" if "requested qty cannot be greater than purchase qty" in lowered else "error"
        payload = {
            "status": response_status,
            "message": message,
            "errors": errors,
        }
        if response_status == "warning":
            payload["requires_confirmation"] = True
        return None, Response(
            payload,
            status=status.HTTP_400_BAD_REQUEST,
        )
    except Exception as exc:
        if isinstance(exc, IntegrityError):
            message = str(exc)
            lowered = message.lower()
            if "quotation_number" in lowered and "item_code" in lowered:
                return None, Response(
                    {
                        "status": "error",
                        "message": "Duplicate item code is not allowed for this quotation.",
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )
        logger.exception("%s", fallback_message)
        return None, Response(
            {
                "status": "error",
                "message": str(exc) or fallback_message,
            },
            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
        )


@api_view(["GET", "POST"])
@csrf_protect
def quotations_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    if request.method == "GET":
        project_id = request.GET.get("project_id")
        payload = ProjectQuotationSummaryView(request).list_payload(project_id=project_id)
        return Response(payload, status=status.HTTP_200_OK)

    serializer = ProjectQuotationSummarySerializer(data=request.data)
    if not serializer.is_valid():
        return _serializer_error_response(serializer)

    summary, error_response = _safe_serializer_save(serializer, "Failed to save quotation summary.")
    if error_response:
        return error_response
    payload = ProjectQuotationSummaryView(request).detail_payload(summary)
    return Response({"success": True, **payload}, status=status.HTTP_201_CREATED)


@api_view(["GET", "PATCH", "DELETE"])
@csrf_protect
def quotation_detail_api_view(request, pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        summary = ProjectQuotationSummaryInfo.objects.select_related("project").get(pk=pk)
    except ProjectQuotationSummaryInfo.DoesNotExist:
        return JsonResponse({"status": "error", "message": "Quotation summary not found."}, status=404)

    if request.method == "GET":
        payload = ProjectQuotationSummaryView(request).detail_payload(summary)
        return Response(payload, status=status.HTTP_200_OK)

    if request.method == "DELETE":
        summary.delete()
        payload = ProjectQuotationSummaryView(request).list_payload()
        return Response({"success": True, "message": "Quotation deleted.", **payload}, status=status.HTTP_200_OK)

    serializer = ProjectQuotationSummarySerializer(summary, data=request.data, partial=True)
    if not serializer.is_valid():
        return _serializer_error_response(serializer)

    updated, error_response = _safe_serializer_save(serializer, "Failed to update quotation summary.")
    if error_response:
        return error_response
    payload = ProjectQuotationSummaryView(request).detail_payload(updated)
    return Response({"success": True, **payload}, status=status.HTTP_200_OK)


@api_view(["GET", "POST"])
@csrf_protect
def quotation_items_api_view(request, pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        summary = ProjectQuotationSummaryInfo.objects.select_related("project").get(pk=pk)
    except ProjectQuotationSummaryInfo.DoesNotExist:
        return JsonResponse({"status": "error", "message": "Quotation summary not found."}, status=404)

    if request.method == "GET":
        payload = ProjectQuotationItemView(request).list_payload(summary)
        return Response(payload, status=status.HTTP_200_OK)

    payload_data = {**request.data, "quotation_number": summary.quotation_number}
    serializer = ProjectQuotationItemSerializer(data=payload_data, context=_item_serializer_context(request))
    if not serializer.is_valid():
        return _serializer_error_response(serializer)

    saved_row, error_response = _safe_serializer_save(serializer, "Failed to save quotation item.")
    if error_response:
        return error_response
    payload = ProjectQuotationSummaryView(request).detail_payload(summary)
    return Response({"success": True, "item_id": saved_row.pk, **payload}, status=status.HTTP_201_CREATED)


@api_view(["PATCH", "DELETE"])
@csrf_protect
def quotation_item_detail_api_view(request, quotation_pk, item_pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        summary = ProjectQuotationSummaryInfo.objects.select_related("project").get(pk=quotation_pk)
    except ProjectQuotationSummaryInfo.DoesNotExist:
        return JsonResponse({"status": "error", "message": "Quotation summary not found."}, status=404)

    try:
        row = ProjectQuotationItemInfo.objects.select_related(
            "quotation_number",
            "cost_type",
            "item_category",
            "item_code__item_type",
            "room_name",
        ).defer(
            "item_type",
        ).get(pk=item_pk, quotation_number=summary)
    except ProjectQuotationItemInfo.DoesNotExist:
        return JsonResponse({"status": "error", "message": "Quotation item not found."}, status=404)

    if request.method == "DELETE":
        delete_pk = row.pk
        row.delete()
        validate_project_quotation_hierarchy(summary, delete_pk=delete_pk)
        summary.save()
        payload = ProjectQuotationSummaryView(request).detail_payload(summary)
        return Response({"success": True, "message": "Quotation item deleted.", **payload}, status=status.HTTP_200_OK)

    serializer = ProjectQuotationItemSerializer(row, data=request.data, partial=True, context=_item_serializer_context(request))
    if not serializer.is_valid():
        return _serializer_error_response(serializer)

    updated_row, error_response = _safe_serializer_save(serializer, "Failed to update quotation item.")
    if error_response:
        return error_response
    payload = ProjectQuotationSummaryView(request).detail_payload(summary)
    return Response({"success": True, "item_id": updated_row.pk, **payload}, status=status.HTTP_200_OK)


@api_view(["GET"])
def list_project_quotations_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    return quotations_api_view(request)


@api_view(["POST"])
@csrf_protect
def create_project_quotation_item_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    serializer = ProjectQuotationItemSerializer(data=request.data, context=_item_serializer_context(request))
    if not serializer.is_valid():
        return _serializer_error_response(serializer)

    row, error_response = _safe_serializer_save(serializer, "Failed to save project quotation item.")
    if error_response:
        return error_response
    payload = ProjectQuotationSummaryView(request).detail_payload(row.quotation_number)
    return Response({"success": True, **payload}, status=status.HTTP_201_CREATED)


@api_view(["GET", "PATCH", "DELETE"])
@csrf_protect
def project_quotation_item_detail_api_view(request, pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        row = ProjectQuotationItemInfo.objects.select_related(
            "quotation_number",
            "cost_type",
            "item_category",
            "item_code__item_type",
            "room_name",
        ).defer(
            "item_type",
        ).get(pk=pk)
    except ProjectQuotationItemInfo.DoesNotExist:
        return JsonResponse({"status": "error", "message": "Project quotation item not found."}, status=404)

    if request.method == "GET":
        return Response(ProjectQuotationSummaryView(request).detail_payload(row.quotation_number), status=status.HTTP_200_OK)

    if request.method == "DELETE":
        summary = row.quotation_number
        delete_pk = row.pk
        row.delete()
        validate_project_quotation_hierarchy(summary, delete_pk=delete_pk)
        summary.save()
        payload = ProjectQuotationSummaryView(request).list_payload()
        return Response({"success": True, "message": "Project quotation item deleted.", **payload}, status=status.HTTP_200_OK)

    serializer = ProjectQuotationItemSerializer(row, data=request.data, partial=True, context=_item_serializer_context(request))
    if not serializer.is_valid():
        return _serializer_error_response(serializer)

    updated_row, error_response = _safe_serializer_save(serializer, "Failed to update project quotation item.")
    if error_response:
        return error_response
    payload = ProjectQuotationSummaryView(request).detail_payload(updated_row.quotation_number)
    return Response({"success": True, **payload}, status=status.HTTP_200_OK)


@api_view(["GET"])
def stock_planning_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    project_id = request.GET.get("project_id")
    payload = ProjectQuotationSummaryView(request).list_payload(project_id=project_id)
    return Response(payload, status=status.HTTP_200_OK)


@api_view(["GET"])
def download_quotation_items_import_template_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    template_bytes = _quotation_items_import_template_bytes()
    if template_bytes is None:
        return Response(
            {"status": "error", "message": "Excel template dependency not installed."},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
        )

    response = HttpResponse(
        template_bytes.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = "attachment; filename=quotation_items_import_template.xlsx"
    return response


@api_view(["POST"])
@csrf_protect
def import_quotation_items_excel_api_view(request, pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        summary = ProjectQuotationSummaryInfo.objects.select_related("project").get(pk=pk)
    except ProjectQuotationSummaryInfo.DoesNotExist:
        return JsonResponse({"status": "error", "message": "Quotation summary not found."}, status=404)

    excel_file = request.FILES.get("file")
    if not excel_file:
        return Response(
            {"status": "error", "message": "Excel file is required (form field: file)."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        import openpyxl
    except ImportError:
        return Response(
            {"status": "error", "message": "Excel import dependency not installed."},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
        )

    try:
        workbook = openpyxl.load_workbook(excel_file, data_only=True)
        worksheet = workbook.active
    except Exception:
        return Response(
            {"status": "error", "message": "Invalid or unreadable Excel file."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return Response(
            {"status": "error", "message": "Excel file is empty."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    column_map = _resolve_import_column_map(rows[0])
    if column_map.get("item_code") is None or column_map.get("requested_qty") is None:
        return Response(
            {
                "status": "error",
                "message": "Excel must include Item Code and Requested Qty columns.",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    material_cost_type = CostTypeInfo.objects.filter(name__iexact="MATERIAL").order_by("id").first()
    if not material_cost_type:
        return Response(
            {
                "status": "error",
                "message": "MATERIAL cost type is not configured.",
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    data_rows = rows[1:] if column_map["has_header"] else rows
    created_count = 0
    updated_count = 0
    failed_count = 0
    blank_count = 0
    row_reports = []

    base_row_number = 2 if column_map["has_header"] else 1
    for index, row_values in enumerate(data_rows, start=base_row_number):
        values = list(row_values or [])

        max_idx = max(i for i in [
            column_map.get("item_category"),
            column_map.get("item_name"),
            column_map.get("item_code"),
            column_map.get("requested_qty"),
            column_map.get("room_name"),
            column_map.get("actual_cost"),
        ] if i is not None)
        if len(values) <= max_idx:
            values.extend([None] * (max_idx + 1 - len(values)))

        item_category_input = _as_excel_text(values[column_map["item_category"]]) if column_map.get("item_category") is not None else ""
        item_name_input = _as_excel_text(values[column_map["item_name"]]) if column_map.get("item_name") is not None else ""
        item_code_input = _as_excel_text(values[column_map["item_code"]])
        requested_qty_input = _as_excel_text(values[column_map["requested_qty"]])
        room_name_input = _as_excel_text(values[column_map["room_name"]]) if column_map.get("room_name") is not None else ""
        actual_cost_input = _as_excel_text(values[column_map["actual_cost"]]) if column_map.get("actual_cost") is not None else ""

        report_payload = {
            "row": index,
            "item_category_input": item_category_input,
            "item_name_input": item_name_input,
            "item_code_input": item_code_input,
            "requested_qty_input": requested_qty_input,
            "room_name_input": room_name_input,
            "actual_cost_input": actual_cost_input,
        }

        if not any(values):
            blank_count += 1
            continue

        if not item_code_input:
            failed_count += 1
            row_reports.append({**report_payload, "status": "failed", "message": "Item Code is required."})
            continue

        if not item_category_input:
            failed_count += 1
            row_reports.append({**report_payload, "status": "failed", "message": "Item Category is required."})
            continue

        if not item_name_input:
            failed_count += 1
            row_reports.append({**report_payload, "status": "failed", "message": "Item Name is required."})
            continue

        warning_notes = []
        requested_qty_value = _to_decimal_or_none(values[column_map["requested_qty"]])
        if requested_qty_value is None:
            requested_qty_value = Decimal("0")
            warning_notes.append("Requested Qty was invalid and defaulted to 0.")
        if requested_qty_value < 0:
            requested_qty_value = Decimal("0")
            warning_notes.append("Requested Qty was negative and defaulted to 0.")

        item_master = (
            LabFurnitureItem.objects.select_related("item_category")
            .filter(item_code__iexact=item_code_input)
            .order_by("id")
            .first()
        )
        if not item_master:
            failed_count += 1
            row_reports.append(
                {
                    **report_payload,
                    "status": "failed",
                    "message": f"Invalid Item Code '{item_code_input}'.",
                }
            )
            continue

        resolved_category_name = _as_excel_text(getattr(getattr(item_master, "item_category", None), "name", ""))
        if item_category_input and item_category_input.lower() != resolved_category_name.lower():
            failed_count += 1
            row_reports.append(
                {
                    **report_payload,
                    "status": "failed",
                    "message": (
                        f"Item Category '{item_category_input}' does not match Item Master "
                        f"('{resolved_category_name or '-'}')."
                    ),
                }
            )
            continue

        resolved_item_name = _as_excel_text(getattr(item_master, "item_name", ""))
        if item_name_input and item_name_input.lower() != resolved_item_name.lower():
            failed_count += 1
            row_reports.append(
                {
                    **report_payload,
                    "status": "failed",
                    "message": (
                        f"Item Name '{item_name_input}' does not match Item Master "
                        f"('{resolved_item_name or '-'}')."
                    ),
                }
            )
            continue

        room_obj = None
        if room_name_input:
            room_obj = RoomDataInfo.objects.filter(room_name__iexact=room_name_input).order_by("id").first()
            if not room_obj:
                warning_notes.append(f"Room '{room_name_input}' not found; room was left blank.")

        existing_row = (
            ProjectQuotationItemInfo.objects.filter(quotation_number=summary, item_code=item_master)
            .order_by("id")
            .first()
        )

        actual_cost_value = _to_decimal_or_none(values[column_map["actual_cost"]]) if column_map.get("actual_cost") is not None else None
        if actual_cost_value is None and existing_row is not None:
            actual_cost_value = existing_row.actual_cost
        if actual_cost_value is None:
            actual_cost_value = Decimal("0")

        payload = {
            "cost_type_id": material_cost_type.pk,
            "item_category_id": item_master.item_category_id,
            "item_name": item_master.item_name,
            "item_code_id": item_master.pk,
            "room_name_id": room_obj.pk if room_obj else None,
            "requested_qty": str(requested_qty_value),
            "actual_cost": str(actual_cost_value),
        }

        context = {"request": request, "allow_requested_qty_override": True}
        if existing_row:
            serializer = ProjectQuotationItemSerializer(existing_row, data=payload, partial=True, context=context)
        else:
            serializer = ProjectQuotationItemSerializer(
                data={**payload, "quotation_number": summary.quotation_number},
                context=context,
            )

        if not serializer.is_valid():
            failed_count += 1
            message = "Validation failed."
            errors = serializer.errors if isinstance(serializer.errors, dict) else {}
            for value in errors.values():
                if isinstance(value, (list, tuple)) and value:
                    message = str(value[0])
                    break
                if isinstance(value, dict):
                    nested_values = list(value.values())
                    if nested_values and isinstance(nested_values[0], (list, tuple)) and nested_values[0]:
                        message = str(nested_values[0][0])
                        break
            row_reports.append({**report_payload, "status": "failed", "message": message})
            continue

        saved_row, error_response = _safe_serializer_save(serializer, "Failed to import quotation item.")
        if error_response:
            failed_count += 1
            payload_data = getattr(error_response, "data", {}) if error_response is not None else {}
            row_reports.append(
                {
                    **report_payload,
                    "status": "failed",
                    "message": str(payload_data.get("message") or "Failed to import row."),
                }
            )
            continue

        if existing_row:
            updated_count += 1
            note_suffix = f" Warnings: {' '.join(warning_notes)}" if warning_notes else ""
            row_reports.append(
                {
                    **report_payload,
                    "status": "updated",
                    "message": f"Updated item {getattr(saved_row.item_code, 'item_code', item_code_input)}.{note_suffix}",
                }
            )
        else:
            created_count += 1
            note_suffix = f" Warnings: {' '.join(warning_notes)}" if warning_notes else ""
            row_reports.append(
                {
                    **report_payload,
                    "status": "created",
                    "message": f"Created item {getattr(saved_row.item_code, 'item_code', item_code_input)}.{note_suffix}",
                }
            )

    payload = ProjectQuotationSummaryView(request).detail_payload(summary)
    return Response(
        {
            "success": True,
            "status": "success",
            "message": "Quotation item import completed.",
            "summary": {
                "created": created_count,
                "updated": updated_count,
                "blank_rows": blank_count,
                "failed": failed_count,
                "total_processed": created_count + updated_count + failed_count,
            },
            "row_reports": row_reports,
            **payload,
        },
        status=status.HTTP_200_OK,
    )


