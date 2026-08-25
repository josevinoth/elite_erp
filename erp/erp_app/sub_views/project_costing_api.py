import logging
from decimal import Decimal
from io import BytesIO

from django.db import transaction
from django.http import HttpResponse, JsonResponse
from django.views.decorators.csrf import csrf_protect
from rest_framework import status
from rest_framework.decorators import api_view
from rest_framework.response import Response

from ..project_costing_items_serializer import ProjectCostingItemSerializer
from ..project_costing_summary_serializer import ProjectCostingSummarySerializer
from ..sub_models.CostType_mod import CostTypeInfo
from ..sub_models.lab_furniture_item import LabFurnitureItem
from ..sub_models.project_costing_items_mod import ProjectCostingItemInfo
from ..sub_models.project_costing_summary_mod import ProjectCostingSummaryInfo
from ..sub_models.project_quotation_summary_mod import ProjectQuotationSummaryInfo
from ..sub_models.retrieval_status_mod import RetrievalStatusInfo
from ..sub_models.room_data_mod import RoomDataInfo
from ..utils import normalize_text
from .project_costing_view import ProjectCostingItemView, ProjectCostingSummaryView

logger = logging.getLogger(__name__)

ALLOWED_RETRIEVAL_STATUSES = {
    RetrievalStatusInfo.STATUS_NO_ACTION,
    RetrievalStatusInfo.STATUS_ITEM_REQUESTED,
    RetrievalStatusInfo.STATUS_ITEM_SUPPLIED,
    RetrievalStatusInfo.STATUS_ITEM_ACCEPTED,
    RetrievalStatusInfo.STATUS_ITEM_RETURN,
    RetrievalStatusInfo.STATUS_ITEM_RETURN_ACCEPTED,
}


def _ensure_authenticated(request):
    if not request.user.is_authenticated:
        return JsonResponse({"status": "error", "message": "Authentication required."}, status=401)
    return None


def _is_admin_user(user):
    if not user or not user.is_authenticated:
        return False
    if user.is_superuser or user.is_staff:
        return True
    return bool({g.name.strip().lower() for g in user.groups.all()}.intersection({"admin", "super admin", "staff"}))


def _user_team_name(user):
    try:
        profile = user.profile
    except Exception:
        return ""
    return str(getattr(getattr(profile, "team", None), "name", "") or "").strip().lower()


def _can_edit_costing_retrieval(user):
    return _is_admin_user(user) or _user_team_name(user) in {"engineering team", "engineering", "engg team"}


def _can_edit_stock_retrieval(user):
    return _is_admin_user(user) or _user_team_name(user) in {"stock team", "stock", "stores team"}


def _status_obj(name):
    normalized_name = str(name or "").strip()
    canonical_name = next((row for row in ALLOWED_RETRIEVAL_STATUSES if row.lower() == normalized_name.lower()), None)
    if not canonical_name:
        return None
    status_obj = RetrievalStatusInfo.objects.filter(status_name__iexact=canonical_name).first()
    if not status_obj:
        status_obj = RetrievalStatusInfo.objects.create(status_name=canonical_name)
    return status_obj


def _is_valid_status_transition(current_name, target_name, source):
    if not target_name:
        return False
    if current_name and current_name.lower() == target_name.lower():
        return True

    source_transition_map = {
        "costing": {
            RetrievalStatusInfo.STATUS_NO_ACTION,
            RetrievalStatusInfo.STATUS_ITEM_REQUESTED,
            RetrievalStatusInfo.STATUS_ITEM_SUPPLIED,
            RetrievalStatusInfo.STATUS_ITEM_ACCEPTED,
            RetrievalStatusInfo.STATUS_ITEM_RETURN,
        },
        "retrieval": {
            RetrievalStatusInfo.STATUS_ITEM_SUPPLIED,
            RetrievalStatusInfo.STATUS_ITEM_ACCEPTED,
            RetrievalStatusInfo.STATUS_ITEM_RETURN,
        },
        "return": {
            RetrievalStatusInfo.STATUS_ITEM_RETURN,
            RetrievalStatusInfo.STATUS_ITEM_RETURN_ACCEPTED,
        },
    }
    return target_name in source_transition_map.get(source, set())


def _summary_not_found():
    return JsonResponse({"status": "error", "message": "Project costing summary not found."}, status=404)


def _item_not_found():
    return JsonResponse({"status": "error", "message": "Project costing item not found."}, status=404)


def _update_costing_items_from_payload(request, summary, items_payload):
    if items_payload in (None, ""):
        return None
    if not isinstance(items_payload, list):
        return Response({"status": "error", "message": "items must be a list."}, status=400)
    if not (_is_admin_user(request.user) or _can_edit_costing_retrieval(request.user)):
        return Response(
            {"status": "error", "message": "Only engineering team or admin can update costing items."},
            status=403,
        )

    for row in items_payload:
        if not isinstance(row, dict):
            return Response({"status": "error", "message": "Each item update must be an object."}, status=400)

        item_pk = row.get("id") or row.get("item_id")
        if not item_pk:
            return Response({"status": "error", "message": "Each item update requires an id."}, status=400)

        try:
            item = ProjectCostingItemInfo.objects.select_related("retrieval_status", "costing_id").get(
                pk=item_pk,
                costing_id=summary,
            )
        except ProjectCostingItemInfo.DoesNotExist:
            return Response(
                {"status": "error", "message": f"Project costing item {item_pk} not found."},
                status=404,
            )

        if item.retrieval_status and item.retrieval_status.status_name == RetrievalStatusInfo.STATUS_ITEM_ACCEPTED and not _is_admin_user(request.user):
            return Response({"status": "error", "message": "Accepted items are frozen."}, status=403)

        payload = {}
        if "requested_qty" in row:
            payload["requested_qty"] = row.get("requested_qty")

        new_status_name = row.get("retrieval_status_name")
        if new_status_name is None and row.get("retrieval_status_id") is not None:
            status_by_id = RetrievalStatusInfo.objects.filter(pk=row.get("retrieval_status_id")).first()
            new_status_name = status_by_id.status_name if status_by_id else None
        if new_status_name is not None:
            status_obj = _status_obj(new_status_name)
            if not status_obj:
                return Response({"status": "error", "message": "Invalid retrieval status."}, status=400)
            current_name = getattr(item.retrieval_status, "status_name", "")
            if not _is_valid_status_transition(current_name, status_obj.status_name, "costing"):
                return Response(
                    {"status": "error", "message": "Invalid retrieval status transition from Project Costing."},
                    status=400,
                )
            payload["retrieval_status_id"] = status_obj.pk

        if not payload:
            continue

        serializer = ProjectCostingItemSerializer(item, data=payload, partial=True)
        if not serializer.is_valid():
            return Response(
                {"status": "error", "message": "Validation failed.", "errors": serializer.errors},
                status=400,
            )
        serializer.save()

    return None


@api_view(["GET", "POST"])
@csrf_protect
def project_costing_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    if request.method == "GET":
        payload = ProjectCostingSummaryView(request).list_payload()
        return Response(payload, status=status.HTTP_200_OK)

    serializer = ProjectCostingSummarySerializer(data=request.data)
    if not serializer.is_valid():
        return Response({"status": "error", "message": "Validation failed.", "errors": serializer.errors}, status=400)

    summary = serializer.save()
    payload = ProjectCostingSummaryView(request).detail_payload(summary)
    return Response({"success": True, **payload}, status=status.HTTP_201_CREATED)


@api_view(["GET", "PATCH", "DELETE"])
@csrf_protect
def project_costing_detail_api_view(request, pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        summary = ProjectCostingSummaryInfo.objects.select_related("project", "quotation_number").get(pk=pk)
    except ProjectCostingSummaryInfo.DoesNotExist:
        return _summary_not_found()

    if request.method == "GET":
        payload = ProjectCostingSummaryView(request).detail_payload(summary)
        return Response(payload, status=status.HTTP_200_OK)

    if request.method == "DELETE":
        summary.delete()
        payload = ProjectCostingSummaryView(request).list_payload()
        return Response({"success": True, "message": "Project costing deleted.", **payload}, status=status.HTTP_200_OK)

    serializer = ProjectCostingSummarySerializer(summary, data=request.data, partial=True)
    if not serializer.is_valid():
        return Response({"status": "error", "message": "Validation failed.", "errors": serializer.errors}, status=400)

    updated = serializer.save()
    payload = ProjectCostingSummaryView(request).detail_payload(updated)
    return Response({"success": True, **payload}, status=status.HTTP_200_OK)


@api_view(["GET", "PATCH"])
@csrf_protect
def project_costing_edit_api_view(request, pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        summary = ProjectCostingSummaryInfo.objects.select_related("project", "quotation_number").get(pk=pk)
    except ProjectCostingSummaryInfo.DoesNotExist:
        return _summary_not_found()

    if request.method == "GET":
        payload = ProjectCostingSummaryView(request).detail_payload(summary)
        return Response(payload, status=status.HTTP_200_OK)

    request_payload = request.data if isinstance(request.data, dict) else {}
    summary_payload = request_payload.get("summary") if isinstance(request_payload.get("summary"), dict) else {
        key: value for key, value in request_payload.items() if key != "items"
    }
    items_payload = request_payload.get("items")

    with transaction.atomic():
        updated_summary = summary
        if summary_payload:
            serializer = ProjectCostingSummarySerializer(summary, data=summary_payload, partial=True)
            if not serializer.is_valid():
                return Response({"status": "error", "message": "Validation failed.", "errors": serializer.errors}, status=400)
            updated_summary = serializer.save()

        item_error = _update_costing_items_from_payload(request, updated_summary, items_payload)
        if item_error is not None:
            return item_error

    payload = ProjectCostingSummaryView(request).detail_payload(updated_summary)
    return Response(
        {"success": True, "message": "Project costing updated successfully.", **payload},
        status=status.HTTP_200_OK,
    )


@api_view(["POST"])
@csrf_protect
def generate_project_costing_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    quotation_id = request.data.get("quotation_id")
    if not quotation_id:
        return Response({"status": "error", "message": "quotation_id is required."}, status=400)

    try:
        quotation = ProjectQuotationSummaryInfo.objects.select_related("project").get(pk=quotation_id)
    except ProjectQuotationSummaryInfo.DoesNotExist:
        return Response({"status": "error", "message": "Quotation summary not found."}, status=404)

    existing = ProjectCostingSummaryInfo.objects.filter(quotation_number=quotation).first()
    if existing:
        payload = ProjectCostingSummaryView(request).detail_payload(existing)
        return Response({"success": True, "message": "Project costing already exists.", **payload}, status=status.HTTP_200_OK)

    costing = ProjectCostingSummaryView(request).clone_from_quotation(quotation)
    payload = ProjectCostingSummaryView(request).detail_payload(costing)
    return Response({"success": True, "message": "Project costing generated successfully.", **payload}, status=status.HTTP_201_CREATED)


@api_view(["GET", "POST"])
@csrf_protect
def project_costing_items_api_view(request, pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        summary = ProjectCostingSummaryInfo.objects.select_related("project", "quotation_number").get(pk=pk)
    except ProjectCostingSummaryInfo.DoesNotExist:
        return _summary_not_found()

    if request.method == "GET":
        payload = ProjectCostingItemView(request).list_payload(summary)
        return Response(payload, status=status.HTTP_200_OK)

    serializer = ProjectCostingItemSerializer(data={**request.data, "costing_id": summary.pk})
    if not serializer.is_valid():
        return Response({"status": "error", "message": "Validation failed.", "errors": serializer.errors}, status=400)

    item = serializer.save()
    payload = ProjectCostingItemView(request).list_payload(summary)
    return Response({"success": True, "item_id": item.pk, **payload}, status=status.HTTP_201_CREATED)


@api_view(["GET", "PATCH", "DELETE"])
@csrf_protect
def project_costing_item_detail_api_view(request, costing_pk, item_pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        summary = ProjectCostingSummaryInfo.objects.select_related("project", "quotation_number").get(pk=costing_pk)
    except ProjectCostingSummaryInfo.DoesNotExist:
        return _summary_not_found()

    try:
        item = ProjectCostingItemInfo.objects.select_related("item_category", "item_code__item_type", "stock_status", "retrieval_status", "costing_id").get(pk=item_pk, costing_id=summary)
    except ProjectCostingItemInfo.DoesNotExist:
        return _item_not_found()

    if request.method == "GET":
        payload = ProjectCostingItemSerializer(item).data
        return Response(payload, status=status.HTTP_200_OK)

    if request.method == "DELETE":
        if not (_is_admin_user(request.user) or _can_edit_costing_retrieval(request.user)):
            return Response({"status": "error", "message": "Only engineering team or admin can delete costing items."}, status=403)
        if item.retrieval_status and item.retrieval_status.status_name == RetrievalStatusInfo.STATUS_ITEM_ACCEPTED and not _is_admin_user(request.user):
            return Response({"status": "error", "message": "Accepted items can only be deleted by admin."}, status=403)
        item.delete()
        payload = ProjectCostingItemView(request).list_payload(summary)
        return Response({"success": True, "message": "Project costing item deleted.", **payload}, status=status.HTTP_200_OK)

    if item.retrieval_status and item.retrieval_status.status_name == RetrievalStatusInfo.STATUS_ITEM_ACCEPTED and not _is_admin_user(request.user):
        return Response({"status": "error", "message": "Accepted items are frozen."}, status=403)
    if not (_is_admin_user(request.user) or _can_edit_costing_retrieval(request.user)):
        return Response({"status": "error", "message": "Only engineering team or admin can update costing items."}, status=403)

    payload = request.data.copy()

    new_status_name = payload.get("retrieval_status_name")
    if new_status_name is None and request.data.get("retrieval_status_id") is not None:
        status_by_id = RetrievalStatusInfo.objects.filter(pk=payload.get("retrieval_status_id")).first()
        new_status_name = status_by_id.status_name if status_by_id else None
    if new_status_name is not None:
        status_obj = _status_obj(new_status_name)
        if not status_obj:
            return Response({"status": "error", "message": "Invalid retrieval status."}, status=400)
        current_name = getattr(item.retrieval_status, "status_name", "")
        if not _is_valid_status_transition(current_name, status_obj.status_name, "costing"):
            return Response({"status": "error", "message": "Invalid retrieval status transition from Project Costing."}, status=400)
        payload["retrieval_status_id"] = status_obj.pk

    serializer = ProjectCostingItemSerializer(item, data=payload, partial=True)
    if not serializer.is_valid():
        return Response({"status": "error", "message": "Validation failed.", "errors": serializer.errors}, status=400)

    updated = serializer.save()
    payload = ProjectCostingItemView(request).list_payload(summary)
    return Response({"success": True, "item_id": updated.pk, **payload}, status=status.HTTP_200_OK)


@api_view(["GET"])
def stock_retrieval_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed
    payload = ProjectCostingItemView(request).stock_retrieval_payload()
    return Response(payload, status=status.HTTP_200_OK)


@api_view(["PATCH"])
@csrf_protect
def stock_retrieval_item_detail_api_view(request, item_pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed
    try:
        item = ProjectCostingItemInfo.objects.select_related("costing_id", "retrieval_status").get(pk=item_pk)
    except ProjectCostingItemInfo.DoesNotExist:
        return _item_not_found()

    if item.retrieval_status and item.retrieval_status.status_name == RetrievalStatusInfo.STATUS_ITEM_ACCEPTED and not _is_admin_user(request.user):
        return Response({"status": "error", "message": "Accepted items are frozen."}, status=403)
    if not _can_edit_stock_retrieval(request.user):
        return Response({"status": "error", "message": "Only stock team or admin can update retrieval status."}, status=403)

    new_status_name = request.data.get("retrieval_status_name") or request.data.get("retrieval_status")
    if not new_status_name:
        return Response({"status": "error", "message": "retrieval_status_name is required."}, status=400)

    status_obj = _status_obj(new_status_name)
    if not status_obj:
        return Response({"status": "error", "message": "Invalid retrieval status."}, status=400)
    current_name = getattr(item.retrieval_status, "status_name", "")
    if not _is_valid_status_transition(current_name, status_obj.status_name, "retrieval"):
        return Response({"status": "error", "message": "Invalid status transition from Stock Retrieval."}, status=400)
    item.retrieval_status = status_obj
    item.save()
    return Response({"success": True, "item": ProjectCostingItemSerializer(item).data}, status=status.HTTP_200_OK)


@api_view(["GET"])
def stock_return_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed
    payload = ProjectCostingItemView(request).stock_return_payload()
    return Response(payload, status=status.HTTP_200_OK)


@api_view(["PATCH"])
@csrf_protect
def stock_return_item_detail_api_view(request, item_pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed
    try:
        item = ProjectCostingItemInfo.objects.select_related("costing_id", "retrieval_status").get(pk=item_pk)
    except ProjectCostingItemInfo.DoesNotExist:
        return _item_not_found()

    if item.retrieval_status and item.retrieval_status.status_name == RetrievalStatusInfo.STATUS_ITEM_ACCEPTED and not _is_admin_user(request.user):
        return Response({"status": "error", "message": "Accepted items are frozen."}, status=403)
    if not _can_edit_stock_retrieval(request.user):
        return Response({"status": "error", "message": "Only stock team or admin can update return status."}, status=403)

    new_status_name = request.data.get("retrieval_status_name") or request.data.get("retrieval_status")
    if not new_status_name:
        return Response({"status": "error", "message": "retrieval_status_name is required."}, status=400)

    status_obj = _status_obj(new_status_name)
    if not status_obj:
        return Response({"status": "error", "message": "Invalid retrieval status."}, status=400)
    current_name = getattr(item.retrieval_status, "status_name", "")
    if not _is_valid_status_transition(current_name, status_obj.status_name, "return"):
        return Response({"status": "error", "message": "Invalid status transition from Stock Return."}, status=400)
    item.retrieval_status = status_obj
    item.save()
    return Response({"success": True, "item": ProjectCostingItemSerializer(item).data}, status=status.HTTP_200_OK)


def _costing_items_import_template_bytes():
    try:
        import openpyxl
    except ImportError:
        return None

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Costing Items Import"
    worksheet.append(["Room Name", "Item Category", "Item Name", "Item Code", "Requested Qty"])
    worksheet.append(["Laboratory", "Category A", "Sample Item", "ITEM01", "1"])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _as_excel_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _to_decimal_or_none(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).strip())
    except Exception:
        return None


@api_view(["GET"])
def download_costing_items_import_template_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    template_bytes = _costing_items_import_template_bytes()
    if template_bytes is None:
        return Response(
            {"status": "error", "message": "Excel template dependency not installed."},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR,
        )

    resp = HttpResponse(
        template_bytes.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = "attachment; filename=costing_items_import_template.xlsx"
    return resp


@api_view(["POST"])
@csrf_protect
def import_costing_items_excel_api_view(request, pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        costing = ProjectCostingSummaryInfo.objects.select_related("project", "quotation_number").get(pk=pk)
    except ProjectCostingSummaryInfo.DoesNotExist:
        return _summary_not_found()

    excel_file = request.FILES.get("file")
    if not excel_file:
        return Response({"status": "error", "message": "Excel file is required."}, status=400)

    try:
        import openpyxl
    except ImportError:
        return Response({"status": "error", "message": "Excel import dependency not installed."}, status=500)

    try:
        workbook = openpyxl.load_workbook(excel_file, data_only=True)
        worksheet = workbook.active
    except Exception:
        return Response({"status": "error", "message": "Invalid or unreadable Excel file."}, status=400)

    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return Response({"status": "error", "message": "Excel file is empty."}, status=400)

    # Detect header row
    first_vals = [normalize_text(v).lower().replace("_", " ") for v in (rows[0] or [])]
    has_header = "item code" in first_vals
    header_map = {h: i for i, h in enumerate(first_vals)} if has_header else {}

    col_room = header_map.get("room name", 0)
    col_category = header_map.get("item category", 1)
    col_name = header_map.get("item name", 2)
    col_code = header_map.get("item code", 3)
    col_qty = header_map.get("requested qty", 4)

    material_ct = CostTypeInfo.objects.filter(name__iexact="MATERIAL").order_by("id").first()

    default_retrieval = RetrievalStatusInfo.objects.filter(
        status_name__iexact=RetrievalStatusInfo.STATUS_NO_ACTION
    ).first()
    if not default_retrieval:
        default_retrieval = RetrievalStatusInfo.objects.create(status_name=RetrievalStatusInfo.STATUS_NO_ACTION)

    data_rows = rows[1:] if has_header else rows
    created_count = updated_count = failed_count = blank_count = 0
    row_reports = []
    base_row = 2 if has_header else 1

    for idx, row_vals in enumerate(data_rows, start=base_row):
        vals = list(row_vals or [])
        max_col = max(col_room, col_category, col_name, col_code, col_qty)
        if len(vals) <= max_col:
            vals.extend([None] * (max_col + 1 - len(vals)))

        if not any(vals):
            blank_count += 1
            continue

        item_code_input = _as_excel_text(vals[col_code])
        item_category_input = _as_excel_text(vals[col_category])
        item_name_input = _as_excel_text(vals[col_name])
        room_name_input = _as_excel_text(vals[col_room])
        requested_qty_input = _as_excel_text(vals[col_qty])

        report_base = {
            "row": idx,
            "item_code_input": item_code_input,
            "item_name_input": item_name_input,
            "item_category_input": item_category_input,
            "room_name_input": room_name_input,
        }

        if not item_code_input:
            failed_count += 1
            row_reports.append({**report_base, "status": "failed", "message": "Item Code is required."})
            continue

        item_master = LabFurnitureItem.objects.select_related("item_category").filter(
            item_code__iexact=item_code_input
        ).order_by("id").first()
        if not item_master:
            failed_count += 1
            row_reports.append({**report_base, "status": "failed", "message": f"Invalid Item Code '{item_code_input}'."})
            continue

        # Validate category if provided
        if item_category_input:
            resolved_cat = _as_excel_text(getattr(getattr(item_master, "item_category", None), "name", ""))
            if item_category_input.lower() != resolved_cat.lower():
                failed_count += 1
                row_reports.append({**report_base, "status": "failed",
                                     "message": f"Item Category '{item_category_input}' does not match Item Master ('{resolved_cat}')."})
                continue

        # Validate item_name if provided
        if item_name_input and item_name_input.lower() != _as_excel_text(getattr(item_master, "item_name", "")).lower():
            failed_count += 1
            row_reports.append({**report_base, "status": "failed",
                                 "message": f"Item Name '{item_name_input}' does not match Item Master."})
            continue

        requested_qty_value = _to_decimal_or_none(requested_qty_input) or Decimal("0")
        if requested_qty_value < 0:
            requested_qty_value = Decimal("0")

        room_obj = None
        if room_name_input:
            room_obj = RoomDataInfo.objects.filter(room_name__iexact=room_name_input).first()

        existing = ProjectCostingItemInfo.objects.filter(costing_id=costing, item_code=item_master).first()

        payload = {
            "costing_id": costing.pk,
            "cost_type_id": material_ct.pk if material_ct else None,
            "item_category_id": item_master.item_category_id,
            "item_name": item_master.item_name,
            "item_code_id": item_master.pk,
            "room_name_id": room_obj.pk if room_obj else None,
            "requested_qty": str(requested_qty_value),
            "retrieval_status_id": default_retrieval.pk if (existing is None) else None,
        }
        if existing is not None:
            payload.pop("retrieval_status_id", None)

        if existing:
            ser = ProjectCostingItemSerializer(existing, data=payload, partial=True)
        else:
            ser = ProjectCostingItemSerializer(data=payload)

        if not ser.is_valid():
            failed_count += 1
            msg = next(
                (str(v[0]) for v in ser.errors.values() if isinstance(v, list) and v),
                "Validation failed."
            )
            row_reports.append({**report_base, "status": "failed", "message": msg})
            continue

        try:
            ser.save()
        except Exception as exc:
            failed_count += 1
            row_reports.append({**report_base, "status": "failed", "message": str(exc)})
            continue

        if existing:
            updated_count += 1
            row_reports.append({**report_base, "status": "updated", "message": f"Updated item {item_code_input}."})
        else:
            created_count += 1
            row_reports.append({**report_base, "status": "created", "message": f"Created item {item_code_input}."})

    payload = ProjectCostingSummaryView(request).detail_payload(costing)
    return Response(
        {
            "success": True,
            "status": "success",
            "message": "Costing item import completed.",
            "summary": {
                "created": created_count,
                "updated": updated_count,
                "blank_rows": blank_count,
                "failed": failed_count,
                "total_processed": created_count + updated_count + failed_count,
            },
            "row_reports": row_reports,
            **payload,
        },
        status=status.HTTP_200_OK,
    )

