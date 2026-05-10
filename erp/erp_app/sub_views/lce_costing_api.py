import json
import io
import datetime
from decimal import Decimal, InvalidOperation

from django.http import FileResponse, JsonResponse
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_GET, require_http_methods, require_POST

from ..sub_models.country_currency import CountryCurrency
from ..sub_models.lce_cost_detail import LCECostDetail
from ..sub_models.project import Project
from ..utils import normalize_text


LCE_FORMAT_COST_HEADS = [
    {"name": "EX WORKS MATERIAL COST", "currency": "FOREIGN", "reference_note": ""},
    {"name": "PACKING CHARGES", "currency": "FOREIGN", "reference_note": ""},
    {"name": "DOCUMENTATION", "currency": "FOREIGN", "reference_note": ""},
    {"name": "OTHER CHARGES 1", "currency": "FOREIGN", "reference_note": ""},
    {"name": "OTHER CHARGES 2", "currency": "FOREIGN", "reference_note": ""},
    {"name": "OTHER CHARGES 3", "currency": "FOREIGN", "reference_note": ""},
    {"name": "OTHER CHARGES 4", "currency": "FOREIGN", "reference_note": ""},
    {"name": "TOTAL SUPPLIER PRICE", "currency": "FOREIGN", "reference_note": "FORMULA"},
    {"name": "ADVANCE PAYMENT VALUE", "currency": "FOREIGN", "reference_note": ""},
    {"name": "BANK EXCHANGE RATE", "currency": "RATE", "reference_note": ""},
    {"name": "ADVANCE PAYMENT VALUE (OMR)", "currency": "OMR", "reference_note": "FORMULA"},
    {"name": "BALANCE PAYMENT VALUE", "currency": "FOREIGN", "reference_note": "FORMULA"},
    {"name": "BALANCE PAYMENT VALUE (OMR)", "currency": "OMR", "reference_note": "FORMULA"},
    {"name": "TOTAL SUPPLIER PRICE (OMR)", "currency": "OMR", "reference_note": "FORMULA"},
    {"name": "BANK MUSCAT CHARGE - ADVANCE PAYMENT", "currency": "OMR", "reference_note": ""},
    {"name": "BANK MUSCAT CHARGE - BALANCE PAYMENT", "currency": "OMR", "reference_note": ""},
    {"name": "FREIGHT CHARGE", "currency": "OMR", "reference_note": ""},
    {"name": "CUSTOMS DUTY (OMR)", "currency": "OMR", "reference_note": ""},
    {"name": "OMAN CUSTOMS BOE CHARGE (OMR)", "currency": "OMR", "reference_note": ""},
    {"name": "ROP CUSTOMS INSPECTION CHARGE", "currency": "OMR", "reference_note": ""},
    {"name": "UNLOADING CHARGE @ MUSCAT STORES 1", "currency": "OMR", "reference_note": ""},
    {"name": "UNLOADING CHARGE @ MUSCAT STORES 2", "currency": "OMR", "reference_note": ""},
    {"name": "LOADING CHARGE @ MUSCAT STORES AT THE TIME OF CUSTOMER DELIVERY", "currency": "OMR", "reference_note": ""},
    {"name": "TOTAL", "currency": "OMR", "reference_note": "FORMULA"},
]


LCE_IMPORT_HEADERS = [
    "S.No",
    "Project",
    "Cost Head",
    "Amount",
    "Currency",
    "Quantity",
    "Unit",
    "Reference Note",
    "Remarks",
    "Updated By",
]


def _excel_text(value):
    if value in (None, ""):
        return ""
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value).strip()


def _split_pk_link(value):
    raw = str(value or "").strip()
    if not raw:
        return None, ""
    parts = [p.strip() for p in raw.split("|", 1)]
    if len(parts) != 2 or not parts[0].isdigit():
        return None, raw
    return int(parts[0]), parts[1]


def _cost_head_defaults(cost_head):
    name = normalize_text(cost_head)
    for row in LCE_FORMAT_COST_HEADS:
        if normalize_text(row.get("name", "")).lower() == name.lower():
            return {
                "currency": row.get("currency", ""),
                "reference_note": row.get("reference_note", ""),
            }
    return {"currency": "", "reference_note": ""}


def _resolve_project_from_import(value):
    if value in (None, ""):
        return None

    pk, label = _split_pk_link(value)
    if pk:
        return Project.objects.filter(pk=pk).first()

    raw = normalize_text(label or value)
    if not raw:
        return None

    if str(raw).isdigit():
        by_pk = Project.objects.filter(pk=int(raw)).first()
        if by_pk:
            return by_pk

    by_project_id = Project.objects.filter(project_id__iexact=raw).order_by("-id").first()
    if by_project_id:
        return by_project_id

    if "-" in raw:
        project_no = normalize_text(raw.split("-", 1)[0])
        by_project_no = Project.objects.filter(project_id__iexact=project_no).order_by("-id").first()
        if by_project_no:
            return by_project_no

    return None


def _build_lce_import_template_bytes():
    try:
        import openpyxl
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LCE Costing Import"
    for col_idx, label in enumerate(LCE_IMPORT_HEADERS, start=1):
        ws.cell(row=1, column=col_idx, value=label)

    ws.freeze_panes = "A2"

    lookup = wb.create_sheet(title="Lookup")
    lookup["A1"] = "Project"
    lookup["B1"] = "Cost Head"

    project_choices = [
        f"{project.pk}|{project.project_id} - {project.project_name}".strip(" -")
        for project in Project.objects.order_by("project_id", "project_name")
    ]

    cost_head_choices = [
        normalize_text(row.get("name", ""))
        for row in LCE_FORMAT_COST_HEADS
        if normalize_text(row.get("name", ""))
    ]
    cost_head_choices = list(dict.fromkeys(cost_head_choices))

    for row_idx, val in enumerate(project_choices, start=1):
        lookup.cell(row=row_idx, column=1, value=val)
    for row_idx, val in enumerate(cost_head_choices, start=1):
        lookup.cell(row=row_idx, column=2, value=val)

    max_input_rows = 5000
    if project_choices:
        dv_project = DataValidation(
            type="list",
            formula1=f"='Lookup'!$A$1:$A${len(project_choices)}",
            allow_blank=True,
        )
        ws.add_data_validation(dv_project)
        dv_project.add(f"B2:B{max_input_rows}")

    if cost_head_choices:
        dv_head = DataValidation(
            type="list",
            formula1=f"='Lookup'!$B$1:$B${len(cost_head_choices)}",
            allow_blank=False,
        )
        ws.add_data_validation(dv_head)
        dv_head.add(f"C2:C{max_input_rows}")

    lookup.sheet_state = "hidden"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _build_lce_export_bytes(rows):
    try:
        import openpyxl
    except ImportError:
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LCE Costing"

    headers = [
        "ID",
        "Project ID",
        "Project Name",
        "Cost Head",
        "Amount",
        "Currency",
        "Quantity",
        "Unit",
        "Line Total",
        "Reference Note",
        "Remarks",
        "Updated By",
    ]

    for col_idx, label in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=label)

    for row_idx, obj in enumerate(rows, start=2):
        line_total = (obj.amount or Decimal("0")) * (obj.quantity or Decimal("0"))
        ws.cell(row=row_idx, column=1, value=obj.pk)
        ws.cell(row=row_idx, column=2, value=obj.project.project_id if obj.project else "")
        ws.cell(row=row_idx, column=3, value=obj.project.project_name if obj.project else "")
        ws.cell(row=row_idx, column=4, value=obj.cost_head)
        ws.cell(row=row_idx, column=5, value=float(obj.amount or 0))
        ws.cell(row=row_idx, column=6, value=obj.currency.currency_code if obj.currency_id else "")
        ws.cell(row=row_idx, column=7, value=float(obj.quantity or 0))
        ws.cell(row=row_idx, column=8, value=obj.unit)
        ws.cell(row=row_idx, column=9, value=float(line_total or 0))
        ws.cell(row=row_idx, column=10, value=obj.reference_note)
        ws.cell(row=row_idx, column=11, value=obj.remarks)
        ws.cell(row=row_idx, column=12, value=obj.created_by)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _ensure_authenticated(request):
    if not request.user.is_authenticated:
        return JsonResponse({"message": "Authentication required."}, status=401)
    return None


def _read_json(request):
    try:
        return json.loads(request.body.decode("utf-8") or "{}")
    except (TypeError, ValueError):
        return {}


def _to_decimal(value, default=Decimal("0")):
    if value in (None, ""):
        return default
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError, TypeError):
        return default


def _to_float(value, default=0.0):
    if value in (None, ""):
        return default
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return default


def _to_project(project_id):
    if project_id in (None, ""):
        return None
    try:
        return Project.objects.get(id=int(project_id))
    except (Project.DoesNotExist, ValueError, TypeError):
        return None


def _serialize_currency(obj):
    if not obj:
        return None
    return {
        "id": obj.id,
        "country_name": obj.country_name,
        "currency_code": obj.currency_code,
        "currency_name": obj.currency_name,
        "label": f"{obj.country_name} - {obj.currency_code}",
    }


def _resolve_currency(value):
    if value in (None, ""):
        return None

    if isinstance(value, CountryCurrency):
        return value

    if isinstance(value, dict):
        value = value.get("id") or value.get("currency_code") or value.get("value") or ""

    if isinstance(value, int) or (isinstance(value, str) and str(value).isdigit()):
        return CountryCurrency.objects.filter(pk=int(value)).first()

    raw = normalize_text(value).upper()
    if not raw:
        return None

    return (
        CountryCurrency.objects.filter(currency_code__iexact=raw).first()
        or CountryCurrency.objects.filter(currency_name__iexact=raw).first()
    )


def _currency_to_omr_factor(currency, fx_rates, default_foreign_rate):
    code = normalize_text(currency).upper()
    if not code or code == "OMR":
        return 1.0
    if code == "RATE":
        return 1.0
    if code in {"FOREIGN", "FOREIGN CURRENCY"}:
        return max(default_foreign_rate, 0.0)
    return max(_to_float(fx_rates.get(code), default_foreign_rate), 0.0)


def _line_total_omr(cost_row, fx_rates, default_foreign_rate):
    amount = _to_float(cost_row.get("amount"), 0.0)
    quantity = _to_float(cost_row.get("quantity"), 1.0)
    factor = _currency_to_omr_factor(cost_row.get("currency"), fx_rates, default_foreign_rate)
    return max(amount, 0.0) * max(quantity, 0.0) * factor


def _driver_value(product, basis):
    if basis == "qty":
        return max(_to_float(product.get("qty"), 0.0), 0.0)
    if basis == "value":
        return max(_to_float(product.get("base_total_omr"), 0.0), 0.0)
    if basis == "weight":
        return max(_to_float(product.get("weight"), 0.0), 0.0)
    if basis == "area":
        return max(_to_float(product.get("area"), 0.0), 0.0)
    if basis == "equal":
        return 1.0
    return max(_to_float(product.get("base_total_omr"), 0.0), 0.0)


def _serialize(obj):
    line_total = (obj.amount or Decimal("0")) * (obj.quantity or Decimal("0"))
    currency_data = _serialize_currency(obj.currency)
    return {
        "id": obj.id,
        "project": obj.project.id if obj.project else None,
        "project_id": obj.project.project_id if obj.project else "",
        "project_name": obj.project.project_name if obj.project else "",
        "cost_head": obj.cost_head,
        "amount": str(obj.amount),
        "currency": currency_data["currency_code"] if currency_data else "",
        "currency_id": currency_data["id"] if currency_data else None,
        "currency_meta": currency_data,
        "quantity": str(obj.quantity),
        "unit": obj.unit,
        "line_total": str(line_total),
        "reference_note": obj.reference_note,
        "remarks": obj.remarks,
        "created_by": obj.created_by,
    }


@require_GET
def download_lce_costing_template_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    content = _build_lce_import_template_bytes()
    if content is None:
        return JsonResponse({"message": "Excel export dependency not installed."}, status=500)

    return FileResponse(
        content,
        as_attachment=True,
        filename="lce_costing_import_template.xlsx",
    )


@require_GET
def export_lce_costing_excel_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    rows = LCECostDetail.objects.select_related("project", "currency").order_by("-id")
    content = _build_lce_export_bytes(rows)
    if content is None:
        return JsonResponse({"message": "Excel export dependency not installed."}, status=500)

    return FileResponse(
        content,
        as_attachment=True,
        filename="lce_costing_export.xlsx",
    )


@require_GET
def list_lce_costing_meta_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    currencies = [
        _serialize_currency(row)
        for row in CountryCurrency.objects.filter(is_active=True).order_by("sort_order", "country_name")
    ]
    return JsonResponse({"cost_heads": LCE_FORMAT_COST_HEADS, "currencies": currencies})


@require_GET
def list_lce_cost_details_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    query = LCECostDetail.objects.select_related("project", "currency").order_by("id")
    return JsonResponse({"lce_cost_details": [_serialize(row) for row in query]})


@require_GET
def list_lce_cost_details_by_project_api_view(request, project_id):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        project = Project.objects.get(pk=project_id)
    except Project.DoesNotExist:
        return JsonResponse({"lce_cost_details": []})

    rows = LCECostDetail.objects.filter(project=project).select_related("project", "currency").order_by("id")
    return JsonResponse({"lce_cost_details": [_serialize(row) for row in rows]})


@require_POST
@csrf_protect
def bulk_save_lce_cost_details_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    payload = _read_json(request)
    project_id = payload.get("project")
    rows_data = payload.get("rows", [])

    project = _to_project(project_id)

    saved = []
    for row_data in rows_data:
        cost_head = normalize_text(row_data.get("cost_head", ""))
        if not cost_head:
            continue

        record_id = row_data.get("id")

        obj = None
        if record_id:
            try:
                obj = LCECostDetail.objects.get(pk=int(record_id))
            except LCECostDetail.DoesNotExist:
                obj = None

        if obj is None:
            obj = (
                LCECostDetail.objects.filter(
                    project=project,
                    cost_head__iexact=cost_head,
                ).first()
                or LCECostDetail(project=project)
            )

        obj.project = project
        obj.cost_head = cost_head
        obj.amount = _to_decimal(row_data.get("amount", getattr(obj, "amount", Decimal("0"))))
        currency_input = row_data.get("currency_id") if row_data.get("currency_id") not in (None, "") else row_data.get("currency")
        resolved_currency = _resolve_currency(currency_input)
        if currency_input not in (None, "") and resolved_currency is None:
            return JsonResponse({"message": f"Invalid currency for cost head: {cost_head}."}, status=400)
        obj.currency = resolved_currency
        obj.quantity = _to_decimal(
            row_data.get("quantity", getattr(obj, "quantity", Decimal("1"))),
            default=Decimal("1"),
        )
        obj.unit = normalize_text(row_data.get("unit", getattr(obj, "unit", "")))
        obj.reference_note = normalize_text(row_data.get("reference_note", getattr(obj, "reference_note", "")))
        obj.remarks = normalize_text(row_data.get("remarks", getattr(obj, "remarks", "")))
        obj.created_by = normalize_text(
            row_data.get("created_by") or request.user.username
        )
        obj.save()
        saved.append(_serialize(obj))

    return JsonResponse({"success": True, "lce_cost_details": saved})


@require_POST
@csrf_protect
def calculate_lce_cost_index_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    payload = _read_json(request)
    project = _to_project(payload.get("project"))
    products = payload.get("products", [])
    fx_rates = payload.get("fx_rates", {})
    default_foreign_rate = _to_float(payload.get("default_foreign_rate"), 0.0)

    if not isinstance(products, list) or not products:
        return JsonResponse({"message": "At least one product is required."}, status=400)

    if project is None:
        return JsonResponse({"message": "Valid project is required."}, status=400)

    product_rows = []
    product_lookup = {}
    for idx, product in enumerate(products):
        qty = max(_to_float(product.get("qty"), 0.0), 0.0)
        if qty <= 0:
            continue

        product_key = normalize_text(product.get("product_code") or product.get("product_name") or f"P{idx + 1}")
        product_currency = normalize_text(product.get("currency", "OMR")) or "OMR"
        product_base_unit = max(_to_float(product.get("base_unit_cost"), 0.0), 0.0)
        product_rate = _currency_to_omr_factor(product_currency, fx_rates, default_foreign_rate)
        base_unit_omr = product_base_unit * product_rate
        base_total_omr = base_unit_omr * qty

        row = {
            "product_key": product_key,
            "product_code": normalize_text(product.get("product_code")),
            "product_name": normalize_text(product.get("product_name")),
            "qty": qty,
            "base_unit_cost": product_base_unit,
            "base_currency": product_currency,
            "base_unit_omr": base_unit_omr,
            "base_total_omr": base_total_omr,
            "weight": max(_to_float(product.get("weight"), 0.0), 0.0),
            "area": max(_to_float(product.get("area"), 0.0), 0.0),
            "allocated_cost_omr": 0.0,
            "allocated_breakdown": [],
        }
        product_rows.append(row)
        product_lookup[product_key] = row

    if not product_rows:
        return JsonResponse({"message": "All product quantities are zero/invalid."}, status=400)

    raw_cost_rows = payload.get("cost_rows")
    if not isinstance(raw_cost_rows, list) or not raw_cost_rows:
        db_rows = LCECostDetail.objects.filter(project=project).order_by("id")
        raw_cost_rows = [
            {
                "id": row.id,
                "cost_head": row.cost_head,
                "amount": str(row.amount),
                "quantity": str(row.quantity),
                "currency": row.currency.currency_code if row.currency_id else "",
                "reference_note": row.reference_note,
                "allocation_basis": "value",
                "applies_to": ["ALL"],
            }
            for row in db_rows
        ]

    for cost_row in raw_cost_rows:
        cost_total_omr = _line_total_omr(cost_row, fx_rates, default_foreign_rate)
        if cost_total_omr <= 0:
            continue

        basis = normalize_text(cost_row.get("allocation_basis", "value")).lower()
        if basis not in {"qty", "value", "weight", "area", "equal", "direct"}:
            basis = "value"

        applies_to = cost_row.get("applies_to")
        if isinstance(applies_to, list):
            raw_targets = [normalize_text(v).upper() for v in applies_to if normalize_text(v)]
        else:
            raw_targets = []

        if not raw_targets or "ALL" in raw_targets:
            targets = list(product_rows)
        else:
            targets = []
            for row in product_rows:
                row_key = normalize_text(row.get("product_key")).upper()
                row_code = normalize_text(row.get("product_code")).upper()
                row_name = normalize_text(row.get("product_name")).upper()
                if row_key in raw_targets or row_code in raw_targets or row_name in raw_targets:
                    targets.append(row)

        if not targets:
            targets = list(product_rows)

        if basis == "direct" and len(targets) == 1:
            target = targets[0]
            target["allocated_cost_omr"] += cost_total_omr
            target["allocated_breakdown"].append(
                {
                    "cost_head": normalize_text(cost_row.get("cost_head")),
                    "basis": "direct",
                    "allocated_omr": round(cost_total_omr, 6),
                }
            )
            continue

        drivers = [_driver_value(target, basis) for target in targets]
        total_driver = sum(drivers)
        if total_driver <= 0:
            drivers = [1.0 for _ in targets]
            total_driver = float(len(targets))

        for target, driver in zip(targets, drivers):
            share = driver / total_driver if total_driver > 0 else 0.0
            allocated = cost_total_omr * share
            target["allocated_cost_omr"] += allocated
            target["allocated_breakdown"].append(
                {
                    "cost_head": normalize_text(cost_row.get("cost_head")),
                    "basis": basis,
                    "allocated_omr": round(allocated, 6),
                }
            )

    product_results = []
    total_base = 0.0
    total_allocated = 0.0
    total_landed = 0.0
    for row in product_rows:
        landed_total_omr = row["base_total_omr"] + row["allocated_cost_omr"]
        unit_landed_omr = landed_total_omr / row["qty"] if row["qty"] > 0 else 0.0
        index_ratio = unit_landed_omr / row["base_unit_omr"] if row["base_unit_omr"] > 0 else 0.0

        total_base += row["base_total_omr"]
        total_allocated += row["allocated_cost_omr"]
        total_landed += landed_total_omr

        product_results.append(
            {
                "product_key": row["product_key"],
                "product_code": row["product_code"],
                "product_name": row["product_name"],
                "qty": round(row["qty"], 6),
                "base_unit_omr": round(row["base_unit_omr"], 6),
                "base_total_omr": round(row["base_total_omr"], 6),
                "allocated_cost_omr": round(row["allocated_cost_omr"], 6),
                "landed_total_omr": round(landed_total_omr, 6),
                "unit_landed_omr": round(unit_landed_omr, 6),
                "cost_index_ratio": round(index_ratio, 6),
                "cost_index_pct": round((index_ratio - 1.0) * 100.0, 3),
                "breakdown": row["allocated_breakdown"],
            }
        )

    return JsonResponse(
        {
            "success": True,
            "project_id": project.project_id,
            "project_name": project.project_name,
            "summary": {
                "total_base_omr": round(total_base, 6),
                "total_allocated_omr": round(total_allocated, 6),
                "total_landed_omr": round(total_landed, 6),
            },
            "products": product_results,
        }
    )


@require_POST
@csrf_protect
def create_lce_cost_detail_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    payload = _read_json(request)
    cost_head = normalize_text(payload.get("cost_head", ""))
    if not cost_head:
        return JsonResponse({"message": "Cost head is required."}, status=400)

    project = _to_project(payload.get("project"))
    currency_input = payload.get("currency_id") if payload.get("currency_id") not in (None, "") else payload.get("currency")
    currency_obj = _resolve_currency(currency_input)

    obj = LCECostDetail.objects.create(
        project=project,
        cost_head=cost_head,
        amount=_to_decimal(payload.get("amount")),
        currency=currency_obj,
        quantity=_to_decimal(payload.get("quantity"), default=Decimal("1")),
        unit=normalize_text(payload.get("unit", "")),
        reference_note=normalize_text(payload.get("reference_note", "")),
        remarks=normalize_text(payload.get("remarks", "")),
        created_by=normalize_text(payload.get("created_by") or request.user.username),
    )
    return JsonResponse({"success": True, "lce_cost_detail": _serialize(obj)}, status=201)


@require_POST
@csrf_protect
def import_lce_costing_excel_api_view(request):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    excel_file = request.FILES.get("file")
    if not excel_file:
        return JsonResponse({"message": "Excel file is required (form field: file)."}, status=400)

    try:
        import openpyxl
    except ImportError:
        return JsonResponse({"message": "Excel import dependency not installed."}, status=500)

    try:
        wb = openpyxl.load_workbook(excel_file, data_only=True)
        ws = wb.active
    except Exception:
        return JsonResponse({"message": "Invalid or unreadable Excel file."}, status=400)

    created_count = 0
    blank_count = 0
    duplicate_count = 0
    failed_count = 0
    row_reports = []

    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = list(row)
        if len(values) < 10:
            values.extend([None] * (10 - len(values)))

        project_raw = values[1]
        cost_head_raw = values[2]
        amount_raw = values[3]
        currency_raw = values[4]
        quantity_raw = values[5]
        unit_raw = values[6]
        reference_note_raw = values[7]
        remarks_raw = values[8]
        created_by_raw = values[9]

        report_row_payload = {
            "row": row_number,
            "project_input": _excel_text(project_raw),
            "cost_head_input": _excel_text(cost_head_raw),
            "amount_input": _excel_text(amount_raw),
            "currency_input": _excel_text(currency_raw),
            "quantity_input": _excel_text(quantity_raw),
            "unit_input": _excel_text(unit_raw),
            "reference_note_input": _excel_text(reference_note_raw),
        }

        if not any([project_raw, cost_head_raw, amount_raw, currency_raw, quantity_raw, unit_raw, reference_note_raw, remarks_raw, created_by_raw]):
            blank_count += 1
            continue

        cost_head = normalize_text(cost_head_raw)
        if not cost_head:
            failed_count += 1
            row_reports.append({**report_row_payload, "status": "failed", "message": "Cost Head is required."})
            continue

        defaults = _cost_head_defaults(cost_head)
        amount = _to_decimal(amount_raw)
        quantity = _to_decimal(quantity_raw, default=Decimal("1"))
        project = _resolve_project_from_import(project_raw)
        currency_code = normalize_text(currency_raw) or defaults.get("currency", "")
        currency_obj = _resolve_currency(currency_code)
        reference_note = normalize_text(reference_note_raw) or defaults.get("reference_note", "")

        if amount < 0:
            failed_count += 1
            row_reports.append({**report_row_payload, "status": "failed", "message": "Amount cannot be negative."})
            continue

        if quantity <= 0:
            failed_count += 1
            row_reports.append({**report_row_payload, "status": "failed", "message": "Quantity must be greater than 0."})
            continue

        exists = LCECostDetail.objects.filter(
            project=project,
            cost_head__iexact=cost_head,
            amount=amount,
            quantity=quantity,
            reference_note__iexact=reference_note,
        ).exists()
        if exists:
            duplicate_count += 1
            row_reports.append(
                {**report_row_payload, "status": "duplicate", "message": "Skipped duplicate row (Project + Cost Head + Amount + Quantity + Reference)."}
            )
            continue

        try:
            obj = LCECostDetail.objects.create(
                project=project,
                cost_head=cost_head,
                amount=amount,
                currency=currency_obj,
                quantity=quantity,
                unit=normalize_text(unit_raw),
                reference_note=reference_note,
                remarks=normalize_text(remarks_raw),
                created_by=normalize_text(created_by_raw) or request.user.username,
            )
            created_count += 1
            row_reports.append(
                {
                    **report_row_payload,
                    "status": "created",
                    "message": f"Imported successfully (ID {obj.pk}).",
                }
            )
        except Exception as exc:
            failed_count += 1
            row_reports.append({**report_row_payload, "status": "failed", "message": str(exc)})

    return JsonResponse(
        {
            "success": True,
            "message": "LCE costing import completed.",
            "summary": {
                "created": created_count,
                "blank_rows": blank_count,
                "duplicates": duplicate_count,
                "failed": failed_count,
            },
            "row_reports": row_reports,
        }
    )


@require_http_methods(["PATCH", "DELETE"])
@csrf_protect
def lce_cost_detail_api_view(request, pk):
    not_allowed = _ensure_authenticated(request)
    if not_allowed:
        return not_allowed

    try:
        obj = LCECostDetail.objects.select_related("project", "currency").get(id=pk)
    except LCECostDetail.DoesNotExist:
        return JsonResponse({"message": "Record not found."}, status=404)

    if request.method == "DELETE":
        obj.delete()
        return JsonResponse({"success": True, "message": "Deleted."})

    payload = _read_json(request)
    cost_head = normalize_text(payload.get("cost_head", obj.cost_head))
    if not cost_head:
        return JsonResponse({"message": "Cost head is required."}, status=400)

    if "project" in payload:
        obj.project = _to_project(payload.get("project"))

    obj.cost_head = cost_head
    obj.amount = _to_decimal(payload.get("amount", obj.amount), default=obj.amount)
    if "currency_id" in payload or "currency" in payload:
        currency_input = payload.get("currency_id") if payload.get("currency_id") not in (None, "") else payload.get("currency")
        obj.currency = _resolve_currency(currency_input)
    obj.quantity = _to_decimal(payload.get("quantity", obj.quantity), default=obj.quantity)
    obj.unit = normalize_text(payload.get("unit", obj.unit))
    obj.reference_note = normalize_text(payload.get("reference_note", obj.reference_note))
    obj.remarks = normalize_text(payload.get("remarks", obj.remarks))
    obj.created_by = normalize_text(payload.get("created_by", obj.created_by))
    obj.save()
    return JsonResponse({"success": True, "lce_cost_detail": _serialize(obj)})

