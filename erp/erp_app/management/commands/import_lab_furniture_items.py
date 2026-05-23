from collections import defaultdict
from pathlib import Path

from django.db.utils import OperationalError, ProgrammingError
from django.core.management.base import BaseCommand, CommandError

from ...sub_models import ItemCategory, LabFurnitureItem
from ...utils import normalize_text


class Command(BaseCommand):
    help = (
        "Import lab furniture items from the DATABASE sheet (column B) of "
        "elite lab furniture - WORK AUTOMATION.xlsx. Bold rows are treated as category headers."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            default="erp_app/media/inputs/elite lab furniture - WORK AUTOMATION.xlsx",
            help="Path to the Excel workbook.",
        )
        parser.add_argument(
            "--sheet",
            default="DATABASE",
            help="Worksheet name to read.",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Preview items without writing to the database.",
        )

    def _load_workbook(self, source_path):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError("openpyxl is required. Install it with: pip install openpyxl") from exc

        try:
            return load_workbook(source_path)
        except FileNotFoundError as exc:
            raise CommandError(f"File not found: {source_path}") from exc
        except Exception as exc:
            raise CommandError(f"Could not open workbook: {exc}") from exc

    def _next_code(self, existing_codes, sequence):
        while True:
            code = f"LF{sequence:04d}"[-6:]
            sequence += 1
            if code not in existing_codes:
                existing_codes.add(code)
                return code, sequence

    def _get_or_create_category(self, name):
        normalized = normalize_text(name)
        if not normalized:
            return None
        category = ItemCategory.objects.filter(name__iexact=normalized).first()
        if category:
            return category
        return ItemCategory.objects.create(name=normalized)

    def handle(self, *args, **options):
        source_path = Path(options["file"])
        sheet_name = options["sheet"]
        dry_run = options["dry_run"]

        wb = self._load_workbook(source_path)
        if sheet_name not in wb.sheetnames:
            raise CommandError(f"Sheet '{sheet_name}' not found. Available: {', '.join(wb.sheetnames)}")

        ws = wb[sheet_name]

        # Example code format before import: LF0001
        self.stdout.write("Example item code format: LF0001 (6-char alphanumeric)")

        try:
            existing_queryset = list(LabFurnitureItem.objects.all())
        except (OperationalError, ProgrammingError):
            existing_queryset = []

        existing_items = {
            (
                normalize_text(item.item_name).casefold(),
                normalize_text(item.item_category.name if item.item_category else "").casefold(),
            ): item
            for item in existing_queryset
        }
        existing_codes = {normalize_text(item.item_code).upper() for item in existing_queryset}
        used_names = defaultdict(int)

        # Find next sequence from existing item codes that look like LF####
        sequence = 1
        for code in existing_codes:
            if code.startswith("LF") and code[2:].isdigit():
                sequence = max(sequence, int(code[2:]) + 1)

        current_category = None
        category_count = 0
        item_count = 0
        preview_rows = []
        seen_in_file = set()

        for row_idx in range(1, ws.max_row + 1):
            cell = ws.cell(row_idx, 2)  # column B
            raw_value = cell.value
            if raw_value in (None, ""):
                continue

            item_text = normalize_text(str(raw_value))
            if not item_text:
                continue

            is_bold = bool(getattr(cell.font, "bold", False))
            if row_idx <= 3:
                continue  # skip title/header rows

            # Bold cells are category headers, except the known sheet title row.
            if is_bold and item_text.upper() != "ITEM DESCRIPTION":
                current_category = self._get_or_create_category(item_text)
                category_count += 1
                continue

            category_name = normalize_text(current_category.name if current_category else "")
            key = (item_text.casefold(), category_name.casefold())
            if key in seen_in_file:
                continue
            seen_in_file.add(key)

            if key in existing_items:
                obj = existing_items[key]
                preview_rows.append(("update", row_idx, category_name, item_text, obj.item_code))
                continue

            code, sequence = self._next_code(existing_codes, sequence)
            used_names[item_text.casefold()] += 1
            preview_rows.append(("create", row_idx, category_name, item_text, code))
            if not dry_run:
                LabFurnitureItem.objects.create(
                    item_name=item_text,
                    item_category=current_category,
                    item_code=code,
                )
            item_count += 1

        self.stdout.write(f"Source file       : {source_path}")
        self.stdout.write(f"Sheet             : {sheet_name}")
        self.stdout.write(f"Categories found  : {category_count}")
        self.stdout.write(f"Items processed   : {len(preview_rows)}")
        self.stdout.write(f"New items         : {item_count}")
        self.stdout.write(f"Dry run           : {'Yes' if dry_run else 'No'}")

        if preview_rows:
            self.stdout.write("\nPreview:")
            for action, row_idx, category, item_name, code in preview_rows[:50]:
                self.stdout.write(f"  {action.upper():6} row {row_idx:>3} | {code} | {category} | {item_name}")
            if len(preview_rows) > 50:
                self.stdout.write(f"  ... and {len(preview_rows) - 50} more rows")

        if dry_run:
            self.stdout.write(self.style.WARNING("\nDry run only. No database changes written."))
            return

        self.stdout.write(self.style.SUCCESS("\nLab furniture item import completed."))


