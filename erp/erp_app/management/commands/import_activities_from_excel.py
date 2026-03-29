from django.core.management.base import BaseCommand, CommandError

from ...sub_models import Activity
from ...utils import normalize_text, to_title_case


class Command(BaseCommand):
    help = (
        "Import unique activities (title-cased) from the Activity column (column F) "
        "of task_details.xlsx into the Activity table."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            default="erp_app/media/task_details.xlsx",
            help="Path to the source Excel file.",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Preview activities to be added without writing to database.",
        )

    def handle(self, *args, **options):
        source_path = options["file"]
        dry_run = options["dry_run"]

        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError(
                "openpyxl is required. Install it with: pip install openpyxl"
            ) from exc

        try:
            wb = load_workbook(source_path, data_only=True)
        except FileNotFoundError as exc:
            raise CommandError(f"File not found: {source_path}") from exc
        except Exception as exc:
            raise CommandError(f"Could not open workbook: {exc}") from exc

        ws = wb.active

        # Collect unique title-cased activities from column F (index 5, 0-based)
        seen = {}  # key: lowercase → value: title-cased name
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw = row[5] if len(row) > 5 else None
            if not raw:
                continue
            name = to_title_case(normalize_text(str(raw)))
            if not name:
                continue
            key = name.lower()
            if key not in seen:
                seen[key] = name

        if not seen:
            self.stdout.write(self.style.WARNING("No activities found in file."))
            return

        # Compare against existing activities
        existing_keys = {
            a.name.lower() for a in Activity.objects.all()
        }

        to_add = [name for key, name in seen.items() if key not in existing_keys]
        already_existing = len(seen) - len(to_add)

        self.stdout.write(f"Source file          : {source_path}")
        self.stdout.write(f"Unique in file       : {len(seen)}")
        self.stdout.write(f"Already in database  : {already_existing}")
        self.stdout.write(f"New to add           : {len(to_add)}")

        if to_add:
            self.stdout.write("\nActivities to add:")
            for name in sorted(to_add):
                self.stdout.write(f"  + {name}")

        if dry_run:
            self.stdout.write(self.style.WARNING("\nDry run only. No changes written."))
            return

        added = 0
        for name in to_add:
            Activity.objects.get_or_create(
                name__iexact=name,
                defaults={"name": name},
            )
            added += 1

        self.stdout.write(self.style.SUCCESS(f"\nActivity import completed. Added: {added}"))

